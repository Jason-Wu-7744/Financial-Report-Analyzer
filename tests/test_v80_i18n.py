"""Offline regression tests: no external APIs or production data."""
import json
import sys
import tempfile
from pathlib import Path
from unittest.mock import patch
from concurrent.futures import ThreadPoolExecutor
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
import i18n
import research_store

@pytest.fixture(autouse=True)
def language_reset(monkeypatch):
    monkeypatch.setenv('DEEPSEEK_API_KEY', '')
    monkeypatch.setenv('OPENAI_API_KEY', '')
    i18n.configure('zh')
    yield
    i18n.configure('zh')


def test_language_is_isolated_between_threads():
    def run(language):
        i18n.configure(language)
        return i18n.translate('登录')
    with ThreadPoolExecutor(2) as pool:
        assert list(pool.map(run,['en','zh'])) == ['Sign in','登录']


def test_catalog_and_numbers():
    i18n.configure('en')
    assert i18n.translate('财报经营分析助手 · V8.0') == 'Financial Report Screening and Analysis System · V8.0'
    assert i18n.translate('2025年半年度报告') == '2025 Interim Report'
    assert i18n.numeric_tokens('收入 123.45，同比增长 +12.3% [S1]') == i18n.numeric_tokens('Revenue 123.45, YoY +12.3% [S1]')
    assert i18n.numeric_tokens('收入 123.45') != i18n.numeric_tokens('Revenue 123.46')
    assert not any(i18n.HAN.search(value) for value in i18n.CATALOG.values())


def test_translation_failure_never_returns_chinese():
    i18n.configure('en')
    with pytest.raises(i18n.TranslationError): i18n.translate('没有词典的动态正文')
    assert not i18n.HAN.search(i18n.translate('没有词典的动态正文',strict=False))


def test_storage_isolates_languages(tmp_path):
    meta=research_store.create_research(tmp_path,'test','Study')
    rid=meta['id']
    research_store.save_summaries(tmp_path,'test',rid,[{'summary':'中文','selected':True}])
    i18n.configure('en')
    assert research_store.load_summaries(tmp_path,'test',rid)==[]
    research_store.save_summaries(tmp_path,'test',rid,[{'summary':'English','selected':False}])
    assert research_store.load_summaries(tmp_path,'test',rid)[0]['summary']=='English'
    i18n.configure('zh')
    assert research_store.load_summaries(tmp_path,'test',rid)[0]['summary']=='中文'


class FixtureTranslator:
    def text(self,value):
        replacements={'示例公司':'Example Company','收入为 123.45，增长 12.3%。':'Revenue was 123.45, up 12.3%.','报告':'Report','证据':'Evidence'}
        for key,translation in replacements.items(): value=value.replace(key,translation)
        assert not i18n.HAN.search(value), repr(value)
        return value


def test_markdown_json_docx_exports_and_originals(tmp_path):
    from docx import Document
    from zipfile import ZipFile
    i18n.configure('en',FixtureTranslator())
    source='# 报告\n\n收入为 123.45，增长 12.3%。'
    md=tmp_path/'original.md'; md.write_text(source,encoding='utf-8')
    result=i18n.localize_artifact(md)
    assert not i18n.HAN.search(result.read_text(encoding='utf-8'))
    assert '123.45' in result.read_text(encoding='utf-8')
    assert md.read_text(encoding='utf-8')==source
    assert i18n.localize_artifact(md)==result
    data=tmp_path/'source.json'; data.write_text(json.dumps({'name':'示例公司','revenue':123.45}),encoding='utf-8')
    converted=json.loads(i18n.localize_artifact(data).read_text(encoding='utf-8'))
    assert converted=={'name':'Example Company','revenue':123.45}
    doc=Document(); doc.add_heading('报告',0); doc.add_paragraph('收入为 123.45，增长 12.3%。')
    doc.add_table(rows=1, cols=1).cell(0,0).text='证据'
    doc.sections[0].header.paragraphs[0].text='示例公司'
    docpath=tmp_path/'source.docx'; doc.save(docpath)
    output=i18n.localize_artifact(docpath)
    with ZipFile(output) as archive:
        for name in archive.namelist():
            if name.startswith('word/') and name.endswith('.xml'):
                from lxml import etree
                tree=etree.fromstring(archive.read(name))
                visible=''.join(tree.xpath('//*[local-name()="t"]/text()'))
                assert not i18n.HAN.search(visible)


def test_generation_keeps_canonical_ids_and_status():
    i18n.configure('en',FixtureTranslator())
    def generate(): return {'company_name':'示例公司','stock_code':'000001.SZ','report_type':'半年度报告','status':'通过','summary':'收入为 123.45，增长 12.3%。','note':'证据','revenue':123.45}
    result=i18n.localized_generation(generate)()
    assert result['company_name']=='示例公司'
    assert result['report_type']=='半年度报告'
    assert result['status']=='通过'
    assert result['revenue']==123.45
    assert result['summary']=='Revenue was 123.45, up 12.3%.'
    assert result['note']=='Evidence'


def test_english_workbook_headers(tmp_path):
    from openpyxl import Workbook
    from engine import read_company_rows
    book=Workbook(); book.active.append(['Ticker','Company']); book.active.append(['600519','Example Company'])
    path=tmp_path/'input.xlsx'; book.save(path)
    _,rows=read_company_rows(path)
    assert rows[0].stock_code=='600519'


def test_cover_login_and_research_english(tmp_path,monkeypatch):
    from streamlit.testing.v1 import AppTest
    import auth
    # All generated files stay in pytest's temporary folder.
    original_user_dir=research_store.user_dir
    monkeypatch.setattr(research_store,'user_dir',lambda root,account: original_user_dir(tmp_path,account))
    monkeypatch.setenv('DEEPSEEK_API_KEY','')
    monkeypatch.setenv('OPENAI_API_KEY','')
    at=AppTest.from_file(str(ROOT/'app.py'), default_timeout=30).run()
    assert not at.exception
    at.selectbox(key='ui_language').set_value('en').run()
    assert not at.exception
    assert [x.label for x in at.text_input]==['Username','Password']
    assert 'V8.1' in '\n'.join(x.value for x in at.markdown)
    at.text_input[0].set_value(auth.configured_accounts()[0][0])
    at.text_input[1].set_value(auth.configured_accounts()[0][1])
    at.button[0].click().run()
    assert not at.exception
    at.button(key='history_new').click().run()
    assert not at.exception
    title=next(x for x in at.text_input if x.label=='Research name')
    title.set_value('Admissions Demo')
    next(x for x in at.button if x.label=='Create').click().run()
    assert not at.exception
    at.run()
    assert 'Research Input' in [x.value for x in at.subheader], (at.session_state.view, [x.label for x in at.button], [x.value for x in at.error])
    assert 'Listing market' in [x.label for x in at.radio]
    for group in ('markdown','caption','info','warning','error','subheader'):
        for element in getattr(at,group):
            assert not i18n.HAN.search(element.value), (group,element.value)
            assert 'translation unavailable' not in element.value.lower(), (group,element.value)
    for group in ('button','text_input','radio','selectbox','number_input','checkbox'):
        for element in getattr(at,group):
            assert not i18n.HAN.search(element.label),(group,element.label)
    # Switch to Chinese and back without losing the project.
    at.selectbox(key='ui_language').set_value('zh').run()
    assert not at.exception
    assert '研究输入' in [x.value for x in at.subheader]
    at.selectbox(key='ui_language').set_value('en').run()
    assert not at.exception
    assert at.session_state.research_id


def test_provider_retries_if_a_figure_changes(monkeypatch):
    from types import SimpleNamespace as NS
    from unittest.mock import Mock
    response=lambda text:NS(choices=[NS(message=NS(content=json.dumps({'text':text})))],usage=NS(prompt_tokens=10,completion_tokens=20))
    create=Mock(side_effect=[response('Revenue 124.45'),response('Revenue __KEEP_A__')])
    monkeypatch.setattr('openai.OpenAI',lambda **kw:NS(chat=NS(completions=NS(create=create))))
    translator=i18n.EnglishTranslator('fake','fake','https://invalid.example')
    assert translator.text('收入 123.45')=='Revenue 123.45'
    assert create.call_count==2
    assert translator.input_tokens==20
    assert translator.text('收入 123.45')=='Revenue 123.45'
    assert create.call_count==2


def test_urls_and_citations_cannot_change():
    assert i18n.protected_tokens('[S1] https://example.com/a') != i18n.protected_tokens('[S2] https://example.com/b')


def test_large_single_paragraph_is_bounded():
    source=('财务信息与证据 123.45。'*2000)
    parts=i18n.bounded_blocks(source)
    assert ''.join(parts)==source
    assert max(map(len,parts))<=8000
    assert sum((i18n.numeric_tokens(part) for part in parts),start=__import__('collections').Counter())==i18n.numeric_tokens(source)


def test_localized_export_is_idempotent(tmp_path):
    i18n.configure('en',FixtureTranslator())
    md=tmp_path/'report.md';md.write_text('# 报告',encoding='utf-8')
    english=i18n.localize_artifact(md)
    assert i18n.localize_artifact(english)==english


def test_failed_summary_does_not_publish_chinese_or_success():
    i18n.configure('en')
    generate=i18n.localized_generation(lambda: {'summary':'无法翻译的实际财务分析','status':'通过','screening_conclusion':'通过'})
    result=generate()
    assert result['translation_failed']
    assert result['status']=='失败'
    assert result['screening_conclusion']=='未完成'
    assert not i18n.HAN.search(result['summary'])


def test_cards_previews_and_demo_chat(tmp_path,monkeypatch):
    from streamlit.testing.v1 import AppTest
    import auth
    original_user_dir=research_store.user_dir
    monkeypatch.setattr(research_store,'user_dir',lambda root,account:original_user_dir(tmp_path,account))
    at=AppTest.from_file(str(ROOT/'app.py'),default_timeout=30)
    at.session_state['ui_language']='en'
    at.session_state['user_name']='Test Reviewer'
    at.session_state['user_account']='guest01'
    at.run()
    i18n.configure('en')
    meta=research_store.create_research(tmp_path,'guest01','Demo')
    at.session_state['research_id']=meta['id']
    at.session_state['view']='research'
    at.session_state['ds_test_mode']=True
    at.session_state['summaries']=[{'stock_code':'600519','company_name':'Example Company','report_type':'半年度报告','report_year':2025,'selected':True,'status':'通过','headline':'Operating growth','driver':'Operating growth','business':'Manufacturing','substance':'Revenue increased by 12.3%.','screening_conclusion':'暂定通过','screening_hits':[],'screening_note':'No verified hard trigger.','revenue_yoy_pct':12.3}]
    at.session_state['llm_drawer_open']=True
    at.run()
    assert not at.exception
    assert 'Quick Analysis结果' not in [x.value for x in at.subheader]
    for group in ('markdown','caption','subheader','info','warning','error'):
        for element in getattr(at,group):
            assert not i18n.HAN.search(element.value),(group,element.value)
            assert 'translation unavailable' not in element.value.lower(),(group,element.value)
    at.chat_input[0].set_value('Explain the earnings drivers.').run()
    assert not at.exception
    for element in at.markdown:
        assert not i18n.HAN.search(element.value)
    # Exercise actual Markdown download controls using an English result fixture.
    md=tmp_path/'fixture_en.md';md.write_text('# Deep Research\n\nRevenue increased 12.3%.',encoding='utf-8')
    at.session_state['comments']=[{'stock_code':'600519','company_name':'Example Company','report_type':'半年度报告','report_year':2025,'path':str(md),'status':'完成','message':'Complete'}]
    at.run()
    assert not at.exception
    assert len(at.get('download_button'))>=1
    for element in at.markdown:
        assert not i18n.HAN.search(element.value)


def test_table_translation_preserves_canonical_values():
    import pandas as pd
    from ui_i18n import StreamlitUI
    captured={}
    class Target:
        def data_editor(self,frame,**kwargs):
            captured.update(frame=frame,config=kwargs['column_config'])
            changed=frame.copy();changed.loc[0,'Research']=False
            return changed
    i18n.configure('en')
    frame=pd.DataFrame([{'研究':True,'证券名称':'示例公司','报告类型':'半年度报告'}])
    returned=StreamlitUI(Target()).data_editor(frame)
    assert captured['frame'].loc[0,'Company']=='Example Company'
    assert captured['config']['Company']=='Company'
    assert returned.loc[0,'证券名称']=='示例公司'
    assert returned.loc[0,'报告类型']=='半年度报告'
    assert not returned.loc[0,'研究']


def test_large_company_list_does_not_call_translation_model():
    import pandas as pd
    from ui_i18n import StreamlitUI
    class FailTranslator:
        cache = {}
        def text(self, value):
            raise AssertionError("Company browsing must not need the model")
    class Target:
        def data_editor(self, frame, **kwargs):
            assert not any(i18n.HAN.search(v) for v in frame['Company'])
            assert frame.iloc[0]['Company'] == 'Gui Zhou Mao Tai'
            return frame
    i18n.configure('en', FailTranslator())
    frame=pd.DataFrame([{'研究':True,'证券代码':str(600000+i),'证券名称':'贵州茅台','报告类型':'半年度报告'} for i in range(803)])
    result=StreamlitUI(Target()).data_editor(frame)
    assert len(result)==803
    assert result.iloc[0]['证券名称']=='贵州茅台'


def test_missing_model_key_error_stays_visible(tmp_path, monkeypatch):
    from streamlit.testing.v1 import AppTest
    original = research_store.user_dir
    monkeypatch.setattr(research_store, 'user_dir', lambda root, account: original(tmp_path, account))
    meta = research_store.create_research(tmp_path, 'reviewer1', 'Missing key validation')
    at = AppTest.from_file(str(ROOT/'app.py'), default_timeout=30)
    at.session_state['ui_language'] = 'en'
    at.session_state['user_name'] = 'Reviewer'
    at.session_state['user_account'] = 'reviewer1'
    at.session_state['research_id'] = meta['id']
    at.session_state['view'] = 'research'
    at.session_state['pending_quick'] = True
    at.session_state['worklist'] = [{'研究': True, '证券代码': '600519', '证券名称': '贵州茅台', '报告类型': '半年度报告', '报告期年份': 2025, '披露日期': '2025-08-13'}]
    at.run()
    assert not at.exception
    assert any('API' in e.value for e in at.error)
    assert not at.session_state['pending_quick']


def test_numeric_validation_ignores_sentence_comma():
    assert i18n.numeric_tokens('收入 123，利润 1,234.56。') == i18n.numeric_tokens('Revenue 123, profit 1,234.56.')
    assert i18n.numeric_tokens('123') != i18n.numeric_tokens('124')


def test_hsbc_formal_interim_title_requires_category():
    from hkex import _is_wanted_report
    title='2025年中期業績報告(附僱員股份計劃)'
    assert _is_wanted_report(title,2025,'半年度报告','財務報表 - [中期/半年度報告]')
    assert not _is_wanted_report(title,2025,'半年度报告','公告及通告 - [中期業績]')
    assert not _is_wanted_report(title,2024,'半年度报告','財務報表 - [中期/半年度報告]')


def test_batch_translation_retries_only_invalid_entries(monkeypatch):
    from types import SimpleNamespace as NS
    from unittest.mock import Mock
    def response(data):
        return NS(choices=[NS(message=NS(content=json.dumps(data)))],usage=None)
    create=Mock(side_effect=[response({'0':'Sales __KEEP_A__','1':'Profit 999'}),response({'0':'Sales __KEEP_A__','1':'Profit 999'}),response({'text':'Profit __KEEP_A__'})])
    monkeypatch.setattr('openai.OpenAI',lambda **kw:NS(chat=NS(completions=NS(create=create))))
    tr=i18n.EnglishTranslator('fake','fake','https://invalid.example')
    tr.prefetch(['销售金额100','利润金额20'])
    assert tr.cache['销售金额100']=='Sales 100'
    assert tr.cache['利润金额20']=='Profit 20'
    assert create.call_count==3


def test_translation_protects_numeric_literals_and_citations():
    source='收入1,234.56，下降-20.1%，见[S12] https://example.com/2025/a.pdf ①'
    masked, tokens=i18n.protected_translation_input(source)
    assert not i18n.numeric_tokens(masked)
    restored=i18n.restore_translation_tokens(masked,tokens)
    assert restored==source
    assert not i18n.restore_translation_tokens(masked.replace('__KEEP_A__',''),tokens)
    assert not i18n.restore_translation_tokens(masked+'__KEEP_A__',tokens)


def test_translation_uses_direct_output_for_deepseek_only():
    assert i18n.translation_api_options('deepseek-v4-flash','https://api.deepseek.com')['extra_body']['thinking']['type']=='disabled'
    assert i18n.translation_api_options('other','https://example.com')=={}


def test_company_html_and_csv_are_english_and_safe():
    import pandas as pd
    from company_list_ui import display_frame, company_table_html
    frame=pd.DataFrame([{'研究':True,'证券代码':'600519','证券名称':'贵州茅台<script>','报告类型':'半年度报告','报告期年份':2025,'披露日期':'2025-08-29'}])
    displayed=display_frame(frame,True)
    assert displayed.iloc[0]['Company']=='贵州茅台<script> (Kweichow Moutai)'
    assert 'Ticker' in displayed.columns and 'Report type' in displayed.columns
    html=company_table_html(displayed)
    assert '<script>' not in html
    assert '600519' in html and '<th scope="col">' in html
    assert frame.iloc[0]['证券名称']=='贵州茅台<script>'

def test_translation_expands_only_new_period_abbreviations():
    masked,tokens=i18n.protected_translation_input('2025年上半年 H1')
    # Original H1 is masked and retained; the model's newly introduced H1 is spelled out.
    translated=masked.replace('年上半年',' first half')+' H1 Q3'
    assert i18n.restore_translation_tokens(translated,tokens)=='2025 first half H1 first half third quarter'


def test_dense_translation_retries_complete_lines(monkeypatch):
    from types import SimpleNamespace as NS
    from unittest.mock import Mock
    def response(text):
        return NS(choices=[NS(message=NS(content=json.dumps({'text':text})))],usage=None)
    create=Mock(side_effect=[response('Missing markers'),response('Missing markers'),response('Revenue __KEEP_A__'),response('Profit __KEEP_A__')])
    monkeypatch.setattr('openai.OpenAI',lambda **kw:NS(chat=NS(completions=NS(create=create))))
    tr=i18n.EnglishTranslator('fake','fake','https://invalid.example')
    assert tr.text('收入100\n利润20')=='Revenue 100\nProfit 20'
    assert create.call_count==4


def test_translation_retries_malformed_json_and_keeps_literal_newline(monkeypatch):
    from types import SimpleNamespace as NS
    from unittest.mock import Mock
    def response(raw):
        return NS(choices=[NS(message=NS(content=raw))],usage=None)
    create=Mock(side_effect=[response('{broken'), response('{"text":"Revenue __KEEP_A__\nverified"}')])
    monkeypatch.setattr('openai.OpenAI',lambda **kw:NS(chat=NS(completions=NS(create=create))))
    tr=i18n.EnglishTranslator('fake','fake','https://invalid.example')
    assert tr.text('收入100\n已核验')=='Revenue 100\nverified'
    assert create.call_count==2


def test_quick_summary_retries_overlong_rewrite_without_slicing(monkeypatch):
    import engine
    from types import SimpleNamespace as NS
    from unittest.mock import Mock
    valid='收入下降但净利润增长，利润增长来自费用控制而非收入扩张，核心产品院内销售为压舱石，院外延伸尚无量数据支撑。经营现金回收仍需跟踪，现有材料未支持业务规模持续扩张的判断。业绩实质：降本驱动的阶段性利润改善，持续性取决于费用压缩空间是否见底，与收入端景气度无关。'
    def response(data):
        return NS(choices=[NS(message=NS(content=json.dumps(data)))],usage=None)
    create=Mock(side_effect=[response({'headline':'费用改善','business':'产品销售','substance_analysis':valid*3}),
                             response({'substance_analysis':valid*3}),response({'substance_analysis':valid})])
    monkeypatch.setattr(engine,'OpenAI',lambda **kw:NS(chat=NS(completions=NS(create=create))))
    result=engine.summarize_with_deepseek_fast(engine.CompanyRow(2,'00700.HK','腾讯控股',{}),[],engine.KeyFinancialMetrics(),'fake')
    assert result[2]==valid
    assert result[-1]==[]
    assert create.call_count==3

def test_new_verbal_classifications_do_not_change_original_figures():
    masked,tokens=i18n.protected_translation_input('一级资本123 第三阶段3')
    translated=masked.replace('一级资本','Tier 1 capital ').replace(' 第三阶段',' Stage 3 ')
    restored=i18n.restore_translation_tokens(translated,tokens)
    assert restored=='Tier one capital 123 Stage three 3'
    assert i18n.numeric_tokens(restored)==i18n.numeric_tokens('一级资本123 第三阶段3')
    assert i18n.restore_translation_tokens('Amount 123',{})=='Amount 123'


def test_checkbox_editor_preserves_hidden_rows_and_canonical_names(monkeypatch):
    import pandas as pd
    import company_list_ui as ui
    class Controls:
        def text_input(self,*args,**kwargs): return '600519'
        def selectbox(self,*args,**kwargs): return 1
        def caption(self,*args,**kwargs): pass
        def download_button(self,*args,**kwargs): pass
    def editor(frame,**kwargs):
        assert len(frame)==1
        assert frame.iloc[0]['Company']=='贵州茅台 (Kweichow Moutai)'
        assert 'Company' in kwargs['disabled']
        assert 'Research' not in kwargs['disabled']
        result=frame.copy(); result['Research']=True
        return result
    monkeypatch.setattr(ui,'st',Controls())
    monkeypatch.setattr(ui.native_st,'data_editor',editor)
    i18n.configure('en')
    frame=pd.DataFrame([
        {'研究':False,'证券代码':'600519','证券名称':'贵州茅台','报告类型':'半年度报告','报告期年份':2025},
        {'研究':True,'证券代码':'300750','证券名称':'宁德时代','报告类型':'半年度报告','报告期年份':2025},
    ])
    result=ui.render_company_list(frame,'test')
    assert result['研究'].tolist()==[True,True]
    assert result['证券名称'].tolist()==['贵州茅台','宁德时代']
    assert not frame.iloc[0]['研究']


def test_quick_rewrite_does_not_copy_long_evidence_into_small_output_budget():
    from engine import CompanyRow, build_substance_rewrite_prompt
    row=CompanyRow(2,'600519','贵州茅台',{})
    evidence='锁定财报原文。'*600
    system,user=build_substance_rewrite_prompt(row,evidence,'过长摘要', ['超过220字'], summary_only=True)
    assert evidence in user
    assert 'evidence_blocks' not in system+user
    assert '只返回 substance_analysis 一个字段' in user
    full_system,full_user=build_substance_rewrite_prompt(row,evidence,'原摘要',[])
    assert 'evidence_blocks' in full_system and '原样返回' in full_user

