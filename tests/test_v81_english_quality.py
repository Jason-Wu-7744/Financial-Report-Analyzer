"""Offline adversarial translation checks for financial meaning and exports."""
import json
import hashlib
import re
from pathlib import Path
import sys
from types import SimpleNamespace as NS
from unittest.mock import Mock

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import i18n


@pytest.fixture(autouse=True)
def language_reset():
    i18n.configure('zh')
    yield
    i18n.configure('zh')


def mock_provider(monkeypatch, outputs):
    def response(value):
        return NS(choices=[NS(message=NS(content=json.dumps(value)))], usage=None)
    create = Mock(side_effect=[response(output) for output in outputs])
    monkeypatch.setattr('openai.OpenAI', lambda **kwargs: NS(chat=NS(completions=NS(create=create))))
    return create


@pytest.mark.parametrize('source, expected', [
    ('收入人民币12.3亿元', 'CNY 12.3 hundred million'),
    ('收入12.3亿港元', 'HKD 12.3 hundred million'),
    ('净亏损美元-1,234.56百万元', 'USD -1,234.56 million'),
    ('现金人民幣−12.5萬元', 'CNY −12.5 ten thousand'),
    ('收入12亿元', '12 hundred million'),
    ('净亏损(12.3亿港元)', '(HKD 12.3 hundred million)'),
])
def test_financial_amount_locks_currency_unit_sign_without_rescaling(source, expected):
    masked, tokens = i18n.protected_translation_input(source, financial_context=True)
    assert expected in tokens.values()
    translated = i18n.restore_translation_tokens(masked.replace('收入', 'Revenue ').replace('净亏损', 'Net loss ').replace('现金', 'Cash '), tokens)
    assert expected in translated
    assert i18n.numeric_tokens(source) == i18n.numeric_tokens(translated)


def test_financial_header_and_uncertainty_are_locked():
    source = '单位：港币百万元。利润可能改善，证据不足，不能排除减值。已核验，未触发。'
    masked, tokens = i18n.protected_translation_input(source, financial_context=True)
    assert {'HKD', 'million', 'may', 'Insufficient evidence', 'Cannot rule out', 'Verified; no trigger identified'} <= set(tokens.values())
    assert not any(word in masked for word in ('港币', '百万元', '可能', '证据不足', '不能排除', '未触发'))


def test_amount_only_labels_are_english_without_api_or_assumed_currency():
    i18n.configure('en')
    assert i18n.translate('12亿元') == '12 hundred million'
    assert i18n.translate(' 12亿港元 ') == ' HKD 12 hundred million '
    assert i18n.EnglishTranslator('', 'fake', 'https://invalid.example').text('12亿港元') == 'HKD 12 hundred million'


def test_english_currency_and_unit_in_chinese_header_cannot_be_rewritten():
    masked, tokens = i18n.protected_translation_input('单位USD million；金额100', financial_context=True)
    assert 'USD' in tokens.values() and 'million' in tokens.values()
    assert 'USD' not in masked and 'million' not in masked


def test_adjacent_currency_and_unit_header_restores_with_english_word_boundary():
    source = '单位：人民币千元'
    masked, tokens = i18n.protected_translation_input(source, financial_context=True)
    restored = i18n.restore_translation_tokens(masked.replace('单位：', 'Unit: '), tokens)
    assert restored == 'Unit: CNY thousand'
    assert not i18n.translation_validation_errors(source, restored)


def test_chinese_date_is_one_locked_fact_without_reordering_numeric_components():
    source = '报告期2025年1月1日—6月30日'
    masked, tokens = i18n.protected_translation_input(source, financial_context=True)
    assert list(tokens.values()) == ['2025/1/1', '6/30']
    translated = i18n.restore_translation_tokens(masked.replace('报告期','Reporting period '), tokens)
    assert not i18n.translation_validation_errors(source, translated)


def test_numeric_ranges_are_locked_to_avoid_double_dashes_or_false_negative_years():
    source = '最近2022-2024年，阈值2至3年，范围1—8页'
    masked,tokens = i18n.protected_translation_input(source,financial_context=True)
    assert list(tokens.values()) == ['2022-2024', '2 to 3', '1—8']
    restored = i18n.restore_translation_tokens(masked.replace('最近','Recent ').replace('年，阈值',' years; threshold ').replace('年，范围',' years; pages ').replace('页',''), tokens)
    assert not i18n.translation_validation_errors(source,restored)


def test_report_terms_are_professional_and_consistent_without_api():
    i18n.configure('en')
    assert i18n.translate('目标公司致命风险排雷报告') == 'Company Critical Risk Screening Report'
    assert i18n.translate('财务真实性存疑') == 'Financial Reporting Reliability Concerns'
    masked,tokens = i18n.protected_translation_input('# 目标公司致命风险排雷报告\n\n判断/口径', financial_context=True)
    assert 'Company Critical Risk Screening Report' in tokens.values()
    assert 'Assessment / Reporting Basis' in tokens.values()


def test_validation_diagnostics_log_only_lengths_and_error_categories(monkeypatch, caplog):
    mock_provider(monkeypatch, [{'text': 'Revenue 124.45'}, {'text': 'Revenue __KEEP_A__'}])
    translator = i18n.EnglishTranslator('private-api-key', 'fake', 'https://invalid.example')
    assert translator.text('收入123.45') == 'Revenue 123.45'
    assert 'source_chars=' in caplog.text and 'errors=' in caplog.text
    assert 'private-api-key' not in caplog.text
    assert '123.45' not in caplog.text and '124.45' not in caplog.text
    assert '收入' not in caplog.text


def test_currency_or_unit_cannot_be_added_outside_intact_markers():
    assert 'currency or unit context changed' in i18n.translation_validation_errors('收入12.3亿港元', 'Revenue CNY HKD 12.3 hundred million')
    assert 'currency or unit context changed' in i18n.translation_validation_errors('金额100百万元', 'Amount 100 billion')
    assert not i18n.translation_validation_errors('金额100百万元', 'Amount 100 million')
    assert not i18n.translation_validation_errors('单位：亿港元', 'Unit: HKD hundred million')


def test_reordered_numbers_or_citations_are_rejected_even_if_counts_match():
    masked, tokens = i18n.protected_translation_input('收入100，利润20，见[S1]')
    reordered = masked.replace('__KEEP_A__', '__SWAP__').replace('__KEEP_B__', '__KEEP_A__').replace('__SWAP__', '__KEEP_B__')
    assert i18n.numeric_tokens('100 20') == i18n.numeric_tokens('20 100')
    assert i18n.restore_translation_tokens(reordered, tokens) == ''


def test_correct_yoy_grammar_can_move_terminology_around_unchanged_percentage():
    source = '同比增长10%'
    masked,tokens=i18n.protected_translation_input(source,financial_context=True)
    assert tokens.fact_keys == ['__KEEP_B__']
    result=i18n.restore_translation_tokens('Growth of __KEEP_B__ __KEEP_A__',tokens)
    assert result == 'Growth of 10% year over year'
    assert not i18n.translation_validation_errors(source,result)


def test_correct_risk_note_can_move_nouns_without_moving_claim_to_another_scope():
    source='未触发致命红线'
    _,tokens=i18n.protected_translation_input(source,financial_context=True)
    result=i18n.restore_translation_tokens('Among __KEEP_B__, there was __KEEP_A__.',tokens)
    assert result == 'Among critical risk thresholds, there was no trigger identified.'
    assert not i18n.translation_validation_errors(source,result)


def test_terminology_cannot_move_to_another_table_cell_or_paragraph():
    source='| 风险一 | 证据不足 |\n| 风险二 | 不适用 |'
    _,tokens=i18n.protected_translation_input(source,financial_context=True)
    assert i18n.restore_translation_tokens('| Risk one | __KEEP_B__ |\n| Risk two | __KEEP_A__ |',tokens) == ''
    _,tokens=i18n.protected_translation_input('证据不足\n\n不适用',financial_context=True)
    assert i18n.restore_translation_tokens('__KEEP_B__\n\n__KEEP_A__',tokens) == ''


def test_marker_adjacent_to_english_word_does_not_hide_currency_or_unit():
    source='金额人民币100百万元'
    _,tokens=i18n.protected_translation_input(source,financial_context=True)
    result=i18n.restore_translation_tokens('Amount __KEEP_A__and growing',tokens)
    assert result == 'Amount CNY 100 million and growing'
    assert not i18n.translation_validation_errors(source,result)


def test_unit_plural_in_header_is_same_unit_and_does_not_become_separate_s():
    source='单位：人民币百万元'
    _,tokens=i18n.protected_translation_input(source,financial_context=True)
    result=i18n.restore_translation_tokens('Unit: __KEEP_A__ __KEEP_B__s',tokens)
    assert result == 'Unit: CNY millions'
    assert not i18n.translation_validation_errors(source,result)
    assert 'currency or unit context changed' in i18n.translation_validation_errors(source,'Unit: CNY billions')
    source='币种人民币'
    _,tokens=i18n.protected_translation_input(source,financial_context=True)
    result=i18n.restore_translation_tokens('Currency__KEEP_A__denominated',tokens)
    assert result == 'Currency CNY denominated'
    assert not i18n.translation_validation_errors(source,result)


def test_accounting_negative_parentheses_are_not_left_for_model_rewriting():
    masked, tokens = i18n.protected_translation_input('净利润(12.3)', financial_context=True)
    assert list(tokens.values()) == ['(12.3)']
    assert '(' not in masked and ')' not in masked


def test_provider_cannot_swap_hkd_for_cny_while_preserving_figure(monkeypatch):
    source = '收入12.3亿港元'
    create = mock_provider(monkeypatch, [{'text': 'Revenue CNY 12.3 hundred million'}, {'text': 'Revenue __KEEP_A__'}])
    translator = i18n.EnglishTranslator('fake', 'fake', 'https://invalid.example')
    assert translator.text(source) == 'Revenue HKD 12.3 hundred million'
    assert create.call_count == 2


def test_anchored_slots_insert_facts_in_source_order_without_model_positioning(monkeypatch):
    create=mock_provider(monkeypatch,[{'segments':['Revenue was ', ', compared with ', '.']}])
    translator=i18n.EnglishTranslator('fake','fake','https://invalid.example')
    result=translator.text('收入100，上期80')
    assert result == 'Revenue was 100, compared with 80.'
    assert not i18n.translation_validation_errors('收入100，上期80',result)
    payload=json.loads(create.call_args.kwargs['messages'][1]['content'])
    assert payload['factual_markers'] == ['__KEEP_A__','__KEEP_B__']


def test_anchored_slots_reject_duplicated_facts_or_new_figures():
    _,tokens=i18n.protected_translation_input('收入100，上期80',financial_context=True)
    assert i18n.translation_response_value({'segments':['__KEEP_B__','','']},tokens) == ''
    reconstructed=i18n.translation_response_value({'segments':['Revenue 999 ', ' versus ', '.']},tokens)
    translated=i18n.restore_translation_tokens(reconstructed,tokens)
    assert 'numeric literals changed' in i18n.translation_validation_errors('收入100，上期80',translated)


def test_pdf_reference_is_one_locked_fact_with_exact_page_list():
    source='见[本期PDF第5、7、28页；上期PDF第4页]'
    _,tokens=i18n.protected_translation_input(source,financial_context=True)
    assert list(tokens.values()) == ['[Current filing PDF pp. 5, 7, 28; Prior filing PDF p. 4]']
    result=i18n.restore_translation_tokens('See __KEEP_A__',tokens)
    assert not i18n.translation_validation_errors(source,result)


@pytest.mark.parametrize('name',['a_deep.md','hk_deep.md','a_risk.md'])
def test_real_canonical_markdown_end_to_end_fact_and_structure_integrity(name,monkeypatch):
    """Actual public-filing QA reports; deterministic backend tests assembly,
    not the linguistic correctness of a provider's translation."""
    source=(ROOT/'tests/fixtures/v81_canonical'/name).read_text(encoding='utf-8-sig')
    calls=[]
    def translated(entry):
        facts=entry['factual_markers']
        semantic=[tag for tag in entry['marker_glossary'] if tag not in facts]
        padding='Reviewed financial narrative. '*max(1,len(i18n.HAN.findall(entry['content']))//12)
        prefix=padding+(' '.join(semantic)+' ' if semantic else '')
        if facts: return {'segments':[prefix]+['; ']*(len(facts)-1)+['.']}
        return prefix+'.'
    def create(**kwargs):
        payload=json.loads(kwargs['messages'][1]['content']);calls.append(payload)
        if 'content' in payload:
            value=translated(payload)
            if isinstance(value,str):value={'text':value}
        else: value={key:translated(entry) for key,entry in payload.items()}
        return NS(choices=[NS(message=NS(content=json.dumps(value)))],usage=None)
    monkeypatch.setattr('openai.OpenAI',lambda **kwargs:NS(chat=NS(completions=NS(create=create))))
    translator=i18n.EnglishTranslator('fake','fake','https://invalid.example')
    result=translator.text(source)
    assert not i18n.translation_validation_errors(source,result)
    assert not i18n.HAN.search(result)
    assert i18n.markdown_structure(source)==i18n.markdown_structure(result)
    assert i18n.numeric_tokens(source)==i18n.numeric_tokens(result)
    assert i18n.financial_context_tokens(source)==i18n.financial_context_tokens(result)
    assert i18n.protected_tokens(source)==i18n.protected_tokens(result)
    count=len(calls)
    assert translator.text(source)==result and len(calls)==count
    assert calls and all(len(entry['factual_markers'])<30 for call in calls for entry in ([call] if 'content' in call else call.values()))


def test_provider_cannot_replace_insufficient_evidence_with_clearance(monkeypatch):
    create = mock_provider(monkeypatch, [{'text': 'The company is safe.'}, {'text': 'Conclusion: __KEEP_A__'}])
    translator = i18n.EnglishTranslator('fake', 'fake', 'https://invalid.example')
    assert translator.text('结论：证据不足') == 'Conclusion: Insufficient evidence'
    assert create.call_count == 2


def test_markdown_collapsed_table_rejected_despite_preserved_figures():
    original = '## 指标\n\n| 指标 | 数值 |\n| --- | --- |\n| 收入 | 100 |\n'
    collapsed = '## Metrics\nRevenue 100'
    assert i18n.numeric_tokens(original) == i18n.numeric_tokens(collapsed)
    assert 'Markdown structure changed' in i18n.translation_validation_errors(original, collapsed)
    proper = '## Metrics\n\n| Metric | Value |\n| --- | --- |\n| Revenue | 100 |\n'
    assert not i18n.translation_validation_errors(original, proper)


def test_materially_truncated_prose_is_rejected_without_numeric_mismatch():
    source = '本期利润改善主要来自成本控制，收入端仍有压力，是否能够持续需要下一期披露进一步验证。'
    assert 'translation appears materially truncated' in i18n.translation_validation_errors(source, 'Good.')


def test_translated_markdown_heading_reused_by_word_export(monkeypatch):
    create = mock_provider(monkeypatch, [{'text': '## Earnings Review\nRevenue __KEEP_A__'}])
    translator = i18n.EnglishTranslator('fake', 'fake', 'https://invalid.example')
    translator.text('## 业绩观察\n收入100')
    assert translator.text('业绩观察') == 'Earnings Review'
    assert translator.text('收入100') == 'Revenue 100'
    assert create.call_count == 1


def test_batch_retry_sends_only_failed_entry(monkeypatch):
    create = mock_provider(monkeypatch, [
        {'0': 'Sales __KEEP_A__', '1': 'Profit 999'},
        {'1': 'Profit __KEEP_A__'},
    ])
    translator = i18n.EnglishTranslator('fake', 'fake', 'https://invalid.example')
    translator.prefetch(['销售金额100', '利润金额20'])
    second_payload = json.loads(create.call_args_list[1].kwargs['messages'][1]['content'])
    assert list(second_payload) == ['1']
    assert translator.cache['销售金额100'] == 'Sales 100'
    assert translator.cache['利润金额20'] == 'Profit 20'


def test_english_export_failure_never_returns_success_or_download(tmp_path, monkeypatch):
    source = tmp_path / 'report.md'
    source.write_text('实际报告内容', encoding='utf-8')
    i18n.configure('en')
    monkeypatch.setattr(i18n, 'localize_artifact', Mock(side_effect=i18n.TranslationError('English validation failed.')))
    result = i18n.localized_generation(lambda: {'status': '完成', 'path': str(source), 'conclusion': '通过'})()
    assert result['status'] == '失败' and result['retryable'] and result['translation_failed']
    assert result['path'] == '' and result['conclusion'] == '未完成'


def test_ui_risk_states_are_distinct_and_no_false_clearance():
    i18n.configure('en')
    values = [i18n.translate(value) for value in ('证据不足', '不适用', '已核验，未触发', '已核验，触发')]
    assert len(set(values)) == 4
    assert 'Pass' not in values
    assert not any(i18n.HAN.search(value) for value in values)


def test_overall_evidence_gap_cannot_be_masked_as_pass():
    masked, tokens = i18n.protected_translation_input('总体结论：暂定待补证', financial_context=True)
    assert list(tokens.values()) == ['Provisional; more evidence needed']
    assert i18n.restore_translation_tokens('Conclusion: Pass', tokens) == ''


def test_failed_localization_keeps_clean_original_for_translation_only_retry(tmp_path, monkeypatch):
    source = tmp_path/'canonical.md'
    source.write_text('# 实际报告', encoding='utf-8')
    english = tmp_path/'english.md'
    english.write_text('# Actual Report', encoding='utf-8')
    original = {'status':'完成', 'path':str(source), 'conclusion':'通过',
                'api_key':'private-key', 'metadata':{'authorization':'Bearer private-key', 'evidence':[1,2]},
                'canonical_result':{'should':'not recurse'}}
    generate = Mock(return_value=original)
    i18n.configure('en')
    monkeypatch.setattr(i18n, 'localize_artifact', Mock(side_effect=i18n.TranslationError('English validation failed.')))
    failed = i18n.localized_generation(generate)()
    canonical = failed['canonical_result']
    assert failed['translation_failed'] and failed['path'] == ''
    assert canonical['path'] == str(source) and canonical['status'] == '完成'
    assert canonical['conclusion'] == '通过' and 'canonical_result' not in canonical
    assert 'api_key' not in canonical and 'authorization' not in canonical['metadata']
    canonical['metadata']['evidence'].append(3)
    assert original['metadata']['evidence'] == [1,2]
    monkeypatch.setattr(i18n, 'localize_artifact', Mock(return_value=english))
    retried = i18n.localized_generation(lambda: canonical)()
    assert retried['status'] == '完成' and retried['path'] == str(english)
    assert not retried.get('translation_failed') and 'canonical_result' not in retried
    assert generate.call_count == 1


def test_summary_translation_failure_keeps_original_chinese_fields():
    i18n.configure('en')
    original = {'status':'通过', 'headline':'原始已验证的财务分析结论', 'substance':'原始完整业务实质内容', 'screening_conclusion':'暂定通过'}
    failed = i18n.localized_generation(lambda: original)()
    assert failed['translation_failed']
    assert failed['canonical_result'] == original
    assert failed['canonical_result']['headline'] != failed['headline']


def test_canonical_retry_state_redacts_credentials_inside_strings():
    translator = i18n.EnglishTranslator('private-secret', 'fake', 'https://invalid.example')
    i18n.configure('en', translator)
    result = i18n.canonical_translation_result({'note':'private-secret was rejected', 'nested':{'DEEPSEEK_API_KEY':'private-secret'}, 'output_tokens':123})
    assert 'private-secret' not in json.dumps(result)
    assert result['output_tokens'] == 123


def test_risk_notes_use_validated_counts_in_natural_english_without_api():
    result = {'review_schema':'evidence-review-v8.1', 'status':'完成', 'conclusion':'待补证', 'retryable':False,
              'hits':[], 'unchecked':['A1','A2'],
              'evidence_counts':{'已核验未触发':1, '已核验触发':0, '证据不足':2, '不适用':1},
              'note':'待补证：2项证据不足', 'message':'待补证：2项证据不足'}
    i18n.configure('en')
    localized = i18n.localize_summary(result)
    assert localized['note'].startswith('More evidence is required for 2 of 4 checks.')
    assert 'overall review remains incomplete' in localized['note']
    assert localized['message'] == localized['note']
    assert localized['conclusion'] == '待补证'
    assert not i18n.HAN.search(localized['note'])
    inconsistent = dict(result, unchecked=['A1'])
    assert i18n.english_risk_note(inconsistent) is None


def test_risk_note_trigger_and_not_applicable_never_imply_safety():
    base = {'review_schema':'evidence-review-v8.1', 'status':'完成', 'retryable':False}
    excluded = i18n.english_risk_note(dict(base,hits=['A1'],unchecked=[],evidence_counts={'已核验未触发':0,'已核验触发':1,'证据不足':0,'不适用':0}))
    assert 'Exclude:' in excluded and 'A1' in excluded
    not_applicable = i18n.english_risk_note(dict(base,hits=[],unchecked=[],evidence_counts={'已核验未触发':0,'已核验触发':0,'证据不足':0,'不适用':19}))
    assert '19 checks were not applicable' in not_applicable
    assert 'Pass' not in not_applicable and 'safe' not in not_applicable


def test_english_json_exports_translate_state_keys_without_mutating_canonical(tmp_path):
    i18n.configure('en')
    source={'evidence_counts':{'已核验未触发':1,'已核验触发':0,'证据不足':18,'不适用':0},
            'canonical_result':{'internal':'不得导出'},'api_key':'must-not-export'}
    path=tmp_path/'risk.json';path.write_text(json.dumps(source,ensure_ascii=False),encoding='utf-8')
    translated=json.loads(i18n.localize_artifact(path).read_text(encoding='utf-8'))
    assert not i18n.HAN.search(json.dumps(translated,ensure_ascii=False))
    assert translated['evidence_counts']['Insufficient evidence']==18
    assert translated['evidence_counts']['Verified; no trigger identified']==1
    assert 'canonical_result' not in translated and 'api_key' not in translated
    assert '证据不足' in source['evidence_counts']
    assert i18n.localize_tree({'证据不足':18})=={'证据不足':18}


@pytest.mark.parametrize('suffix',['json','docx'])
def test_actual_risk_download_body_and_filename_are_english(suffix,tmp_path,monkeypatch):
    from shutil import copyfile
    from zipfile import ZipFile
    from lxml import etree
    source=ROOT/'tests/fixtures/v81_canonical'/f'a_risk.{suffix}'
    path=tmp_path/f'贵州茅台排雷报告.{suffix}'
    copyfile(source,path)
    def translated(entry):
        facts=entry['factual_markers'];semantic=[tag for tag in entry['marker_glossary'] if tag not in facts]
        label=''.join(chr(97+byte%26) for byte in hashlib.sha256(entry['content'].encode()).digest()[:10])
        prefix=('Reviewed financial narrative '+label+'. ')*max(1,len(i18n.HAN.findall(entry['content']))//12)
        prefix+=(' '.join(semantic)+' ' if semantic else '')
        return {'segments':[prefix]+['; ']*(len(facts)-1)+['.']} if facts else prefix+'.'
    def create(**kwargs):
        payload=json.loads(kwargs['messages'][1]['content'])
        if 'content' in payload:
            result=translated(payload)
            if isinstance(result,str):result={'text':result}
        else:result={key:translated(entry) for key,entry in payload.items()}
        return NS(choices=[NS(message=NS(content=json.dumps(result)))],usage=None)
    monkeypatch.setattr('openai.OpenAI',lambda **kwargs:NS(chat=NS(completions=NS(create=create))))
    i18n.configure('en',i18n.EnglishTranslator('fake','fake','https://invalid.example'))
    output=i18n.localize_artifact(path)
    assert output.name.isascii() and output.name.endswith('_en.'+suffix)
    if suffix=='json':
        data=json.loads(output.read_text(encoding='utf-8'))
        assert not i18n.HAN.search(json.dumps(data,ensure_ascii=False))
        assert len(data['checks'])==19
        assert sum(data['evidence_counts'].values())==19
    else:
        with ZipFile(output) as archive:
            for filename in archive.namelist():
                if filename.startswith('word/') and filename.endswith('.xml'):
                    text=''.join(etree.fromstring(archive.read(filename)).xpath('//*[local-name()="t"]/text()'))
                    assert not i18n.HAN.search(text)
        with ZipFile(path) as archive:
            original=''.join(etree.fromstring(archive.read('word/document.xml')).xpath('//*[local-name()="t"]/text()'))
        with ZipFile(output) as archive:
            converted=''.join(etree.fromstring(archive.read('word/document.xml')).xpath('//*[local-name()="t"]/text()'))
        assert i18n.numeric_tokens(original)==i18n.numeric_tokens(converted)


def test_generation_download_filename_is_english_after_localization(tmp_path,monkeypatch):
    original=tmp_path/'贵州茅台报告.md';original.write_text('原文',encoding='utf-8')
    converted=tmp_path/'report_abcdef_en.md';converted.write_text('English',encoding='utf-8')
    i18n.configure('en')
    monkeypatch.setattr(i18n,'localize_artifact',Mock(return_value=converted))
    result=i18n.localized_generation(lambda:{'status':'完成','path':str(original),'filename':original.name})()
    assert result['filename']=='report_abcdef_en.md'


def test_chinese_pdf_citations_and_review_limits_remain_clear():
    i18n.configure('en')
    assert i18n.translate('模型读取及页码定位，待人工复核')=='Model reading and page identification; pending human review'
    assert i18n.translate('已观察到的风险事实')=='Observed risk facts'
    source='见[本期PDF第5页]'
    assert i18n.protected_tokens(source)!=i18n.protected_tokens('See [Current filing PDF p. 6]')


def test_legacy_english_json_cache_is_migrated_without_any_model(tmp_path,monkeypatch):
    directory=tmp_path/'.english';directory.mkdir()
    path=directory/'report_0123456789abcdef0123_en.json'
    source={'evidence_counts':{'已核验未触发':1,'已核验触发':0,'证据不足':18,'不适用':0},'note':'English only.'}
    path.write_text(json.dumps(source,ensure_ascii=False),encoding='utf-8')
    monkeypatch.setattr('openai.OpenAI',Mock(side_effect=AssertionError('No API calls permitted')))
    i18n.configure('en')
    assert i18n.localize_artifact(path)==path
    data=json.loads(path.read_text(encoding='utf-8'))
    assert not i18n.HAN.search(json.dumps(data,ensure_ascii=False))
    assert data['evidence_counts']['Insufficient evidence']==18
    assert path.with_suffix('.json.before_v81_english_repair').exists()
    assert '证据不足' in source['evidence_counts']


def test_legacy_json_unknown_chinese_fails_closed_without_model(tmp_path,monkeypatch):
    directory=tmp_path/'.english';directory.mkdir()
    path=directory/'report_0123456789abcdef0123_en.json'
    path.write_text(json.dumps({'note':'无法通过静态词典翻译的真实财务正文'},ensure_ascii=False),encoding='utf-8')
    original=path.read_bytes()
    monkeypatch.setattr('openai.OpenAI',Mock(side_effect=AssertionError('No API calls permitted')))
    i18n.configure('en')
    with pytest.raises(i18n.TranslationError):i18n.localize_artifact(path)
    assert path.read_bytes()==original


def test_legacy_docx_cache_keeps_english_and_removes_overclaiming_notice(tmp_path,monkeypatch):
    from docx import Document
    directory=tmp_path/'.english';directory.mkdir()
    path=directory/'report_0123456789abcdef0123_en.docx'
    doc=Document();doc.add_paragraph('Revenue 100. English translation of the validated source-language analysis.');doc.save(path)
    monkeypatch.setattr('openai.OpenAI',Mock(side_effect=AssertionError('No API calls permitted')))
    i18n.configure('en')
    i18n.localize_artifact(path)
    body='\n'.join(p.text for p in Document(path).paragraphs)
    assert 'validated' not in body and 'Revenue 100.' in body
    assert not i18n.HAN.search(body)


def test_year_qualified_half_year_spacing_is_not_a_new_risk_reference():
    source='2025H1收入100；2024H2收入90。风险H1、H2；见[本期PDF第5页]。'
    target='2025 H1 revenue 100; 2024 H2 revenue 90. Risks H1, H2; see [Current filing PDF p. 5].'
    assert i18n.protected_tokens(source)==i18n.protected_tokens(target)
    assert i18n.protected_tokens(target)['H1']==1
    assert i18n.protected_tokens(target)['H2']==1
    assert i18n.protected_tokens(source)==i18n.protected_tokens(target.replace('2025 H1','2025, H1'))
    assert 'references changed' not in i18n.translation_validation_errors(source,target)
    assert i18n.protected_tokens(source)!=i18n.protected_tokens(target.replace('Risks H1, H2','Risks H2, H2'))


def test_half_year_reference_disambiguation_does_not_modify_url_integrity():
    source='https://example.test/2025H1/H1.pdf 风险H1'
    assert i18n.protected_tokens(source)['H1']==1
    assert i18n.protected_tokens(source)!=i18n.protected_tokens(source.replace('/2025H1/','/2025H2/'))


def test_primary_text_protocol_ignores_malformed_auxiliary_slots(monkeypatch):
    create=mock_provider(monkeypatch,[{'text':'Revenue __KEEP_A__, prior period __KEEP_B__.',
                                     'segments':['Revenue ','','','']}])
    translator=i18n.EnglishTranslator('fake','fake','https://invalid.example')
    assert translator.text('收入100，上期80')=='Revenue 100, prior period 80.'
    instruction=create.call_args.kwargs['messages'][0]['content']
    assert 'instead return' not in instruction and 'NEVER include any factual marker' not in instruction
    assert 'Copy every supplied marker exactly once into text' in instruction
    assert create.call_count==1


def test_malformed_slots_stop_after_two_batch_attempts_and_keep_checked_work(monkeypatch):
    invalid={'segments':['Profit ','','','']}
    create=mock_provider(monkeypatch,[{'0':'Sales __KEEP_A__','1':invalid},{'1':invalid}])
    translator=i18n.EnglishTranslator('fake','fake','https://invalid.example')
    with pytest.raises(i18n.TranslationError,match='invalid batch content'):
        translator.prefetch(['销售金额100','利润金额20'])
    assert create.call_count==2
    assert translator.cache['销售金额100']=='Sales 100'
    assert '利润金额20' not in translator.cache
    instruction=create.call_args.kwargs['messages'][0]['content']
    assert 'mapped to {"segments"' not in instruction
    assert 'including ALL factual_markers' in instruction


def test_empty_batch_mappings_do_not_fan_out_into_sentence_requests(monkeypatch):
    create=mock_provider(monkeypatch,[{},{}])
    translator=i18n.EnglishTranslator('fake','fake','https://invalid.example')
    with pytest.raises(i18n.TranslationError,match='invalid batch content'):
        translator.prefetch(['销售金额100','利润金额20','现金金额30'])
    assert create.call_count==2 and not translator.cache


@pytest.mark.parametrize('wrapped',[
    {'translations':{'0':{'text':'Sales __KEEP_A__'},'1':{'text':'Profit __KEEP_A__'}}},
    {'translations':[{'id':'0','text':'Sales __KEEP_A__'},{'id':'1','text':'Profit __KEEP_A__'}]},
])
def test_explicit_batch_id_wrappers_are_compatible_and_validated(wrapped,monkeypatch):
    create=mock_provider(monkeypatch,[wrapped])
    translator=i18n.EnglishTranslator('fake','fake','https://invalid.example')
    translator.prefetch(['销售金额100','利润金额20'])
    assert translator.cache['销售金额100']=='Sales 100'
    assert translator.cache['利润金额20']=='Profit 20'
    assert create.call_count==1


def test_batch_compatibility_never_guesses_ids_or_accepts_duplicate_ids():
    assert i18n.translation_batch_items(['Sales','Profit'],['0','1'])=={}
    assert i18n.translation_batch_items([{'id':'0','text':'Sales'},{'id':'0','text':'Profit'}],['0','1'])=={}
    assert i18n.translation_batch_items([{'id':'7','text':'Sales'}],['0','1'])=={}


def test_text_protocol_fact_reordering_remains_fail_closed(monkeypatch):
    create=mock_provider(monkeypatch,[{'text':'Revenue __KEEP_B__, prior period __KEEP_A__.'}]*2)
    translator=i18n.EnglishTranslator('fake','fake','https://invalid.example')
    with pytest.raises(i18n.TranslationError):translator.text('收入100，上期80')
    assert create.call_count==2 and not translator.cache


def test_indexed_legacy_slots_require_every_exact_index():
    _,tokens=i18n.protected_translation_input('收入100，上期80',financial_context=True)
    assert i18n.translation_response_value({'segments':{'0':'Revenue ','1':', prior ','2':'.'}},tokens)=='Revenue __KEEP_A__, prior __KEEP_B__.'
    assert i18n.translation_response_value({'segments':{'0':'Revenue ','2':'.'}},tokens)==''


def test_order_only_failure_recovers_original_prose_spans_with_fixed_facts(monkeypatch):
    source='收入100，上期80'
    reversed_text='Revenue __KEEP_B__, prior __KEEP_A__'
    create=mock_provider(monkeypatch,[{'0':reversed_text},{'0':reversed_text},{'0':'Revenue','2':', prior'}])
    translator=i18n.EnglishTranslator('fake','fake','https://invalid.example')
    translator.prefetch([source])
    assert translator.cache[source]=='Revenue 100, prior 80'
    assert not i18n.translation_validation_errors(source,translator.cache[source])
    assert create.call_count==3 and translator.fixed_fact_fallback_calls==1
    payload=json.loads(create.call_args.kwargs['messages'][1]['content'])
    assert payload['prose_spans']=={'0':'收入','2':'，上期'}
    assert all('__KEEP_' not in span for span in payload['prose_spans'].values())


def test_fixed_fact_recovery_rejects_new_figures_and_keeps_source_unpublished(monkeypatch):
    source='收入100，上期80'
    reversed_text='Revenue __KEEP_B__, prior __KEEP_A__'
    create=mock_provider(monkeypatch,[{'0':reversed_text},{'0':reversed_text},
                                     {'0':'Revenue 999','2':', prior'},{'0':'Revenue 999'}])
    translator=i18n.EnglishTranslator('fake','fake','https://invalid.example')
    with pytest.raises(i18n.TranslationError,match='fixed-fact translation'):
        translator.prefetch([source])
    assert source not in translator.cache and create.call_count==4
    assert translator.fixed_fact_fallback_calls==2


def test_fixed_fact_recovery_has_a_session_wide_call_budget(monkeypatch):
    source='收入100，上期80'
    reversed_text='Revenue __KEEP_B__, prior __KEEP_A__'
    create=mock_provider(monkeypatch,[{'0':reversed_text},{'0':reversed_text}])
    translator=i18n.EnglishTranslator('fake','fake','https://invalid.example')
    translator.fixed_fact_fallback_calls=translator.max_fixed_fact_fallback_calls
    with pytest.raises(i18n.TranslationError,match='recovery limit'):
        translator.prefetch([source])
    assert create.call_count==2 and source not in translator.cache


def test_fixed_fact_fallback_eligibility_excludes_missing_markers_and_moved_cells():
    masked,tokens=i18n.protected_translation_input('| 收入100 | 上期80 |',financial_context=True)
    assert i18n.only_factual_marker_order_changed('| Revenue __KEEP_B__ | Prior __KEEP_A__ |',tokens) is False
    assert i18n.only_factual_marker_order_changed('| Revenue __KEEP_A__ | Prior |',tokens) is False
    plan=i18n.fixed_fact_prose_plan(masked,tokens)
    assert [part for part,needed in plan if part in tokens.fact_keys]==tokens.fact_keys
    assert all(not any(tag in part for tag in tokens.fact_keys) for part,needed in plan if needed)


def test_fixed_fact_recovery_preserves_table_scopes_currency_citations_and_risk(monkeypatch):
    source='| 净利润12亿港元，去年10亿港元 | 证据不足，见[本期PDF第5页] |'
    masked,tokens=i18n.protected_translation_input(source,financial_context=True)
    plan=i18n.fixed_fact_prose_plan(masked,tokens)
    parts={str(i):part for i,(part,needed) in enumerate(plan) if needed}
    values={key:part.replace('净利润','Net profit').replace('，去年',', prior').replace('，见',', see') for key,part in parts.items()}
    create=mock_provider(monkeypatch,[values])
    translator=i18n.EnglishTranslator('fake','fake','https://invalid.example')
    from openai import OpenAI
    result=translator._translate_fixed_fact_prose(source,masked,tokens,OpenAI(api_key='fake'))
    assert not i18n.translation_validation_errors(source,result)
    assert 'HKD 12 hundred million' in result and 'HKD 10 hundred million' in result
    assert 'Insufficient evidence' in result and '[Current filing PDF p. 5]' in result
    assert i18n.markdown_structure(source)==i18n.markdown_structure(result)
    assert create.call_count==1


def test_four_order_only_items_recover_without_discarding_checked_batch(monkeypatch):
    sources=[f'收入{100+i}，上期{80+i}' for i in range(16)]
    reversed_text='Revenue __KEEP_B__, prior __KEEP_A__'
    first={str(i):(reversed_text if i<4 else 'Revenue __KEEP_A__, prior __KEEP_B__') for i in range(16)}
    second={str(i):reversed_text for i in range(4)}
    create=mock_provider(monkeypatch,[first,second]+[{'0':'Revenue','2':', prior'}]*4)
    translator=i18n.EnglishTranslator('fake','fake','https://invalid.example')
    translator.prefetch(sources)
    assert all(source in translator.cache for source in sources)
    assert all(not i18n.translation_validation_errors(source,translator.cache[source]) for source in sources)
    assert create.call_count==6 and translator.fixed_fact_fallback_calls==4
    retry_payload=json.loads(create.call_args_list[1].kwargs['messages'][1]['content'])
    assert set(retry_payload)=={'0','1','2','3'}


def test_fixed_fact_prefetch_budget_is_lower_than_session_budget(monkeypatch):
    source='收入100，上期80'
    reversed_text='Revenue __KEEP_B__, prior __KEEP_A__'
    create=mock_provider(monkeypatch,[{'0':reversed_text},{'0':reversed_text}])
    translator=i18n.EnglishTranslator('fake','fake','https://invalid.example')
    assert translator.max_fixed_fact_fallback_calls==64
    assert translator.max_prefetch_fixed_fact_fallback_calls==32
    translator.max_prefetch_fixed_fact_fallback_calls=0
    with pytest.raises(i18n.TranslationError,match='recovery limit'):
        translator.prefetch([source])
    assert create.call_count==2 and translator.fixed_fact_fallback_calls==0


def test_three_reports_can_recover_eleven_order_failures_in_one_session(monkeypatch):
    calls=[]
    def create(**kwargs):
        payload=json.loads(kwargs['messages'][1]['content']);calls.append(payload)
        if 'prose_spans' in payload:
            output={key:span.replace('收入','Revenue').replace('，上期',', prior') for key,span in payload['prose_spans'].items()}
        else:
            failures=4 if len(payload)>=16 else 3 if len(payload)>=12 else len(payload)
            output={key:('Revenue __KEEP_B__, prior __KEEP_A__' if int(key)<failures else 'Revenue __KEEP_A__, prior __KEEP_B__') for key in payload}
        return NS(choices=[NS(message=NS(content=json.dumps(output)))],usage=None)
    monkeypatch.setattr('openai.OpenAI',lambda **kwargs:NS(chat=NS(completions=NS(create=create))))
    translator=i18n.EnglishTranslator('fake','fake','https://invalid.example')
    for base,size in ((1000,16),(2000,16),(3000,12)):
        sources=[f'收入{base+i}，上期{base-500+i}' for i in range(size)]
        translator.prefetch(sources)
        assert all(not i18n.translation_validation_errors(source,translator.cache[source]) for source in sources)
    assert translator.fixed_fact_fallback_calls==11
    assert len(calls)==17


def test_sixteen_of_sixty_five_order_only_failures_are_recoverable(monkeypatch):
    sources=[f'收入{100+i}，上期{80+i}' for i in range(65)]
    reversed_text='Revenue __KEEP_B__, prior __KEEP_A__'
    first={str(i):(reversed_text if i<16 else 'Revenue __KEEP_A__, prior __KEEP_B__') for i in range(65)}
    second={str(i):reversed_text for i in range(16)}
    create=mock_provider(monkeypatch,[first,second]+[{'0':'Revenue','2':', prior'}]*16)
    translator=i18n.EnglishTranslator('fake','fake','https://invalid.example')
    translator.prefetch(sources)
    assert all(not i18n.translation_validation_errors(source,translator.cache[source]) for source in sources)
    assert create.call_count==18 and translator.fixed_fact_fallback_calls==16


def test_last_prose_span_has_one_single_text_recovery_without_changing_facts(monkeypatch):
    source='收入100，上期80，增长20%，目标90，预测120'
    masked,tokens=i18n.protected_translation_input(source,financial_context=True)
    requested={str(i):part for i,(part,needed) in enumerate(i18n.fixed_fact_prose_plan(masked,tokens)) if needed}
    assert len(requested)==5
    bad_key=list(requested)[-1]
    good={key:part.replace('收入','Revenue').replace('，上期',', prior').replace('，增长',', growth').replace('，目标',', target') for key,part in requested.items() if key!=bad_key}
    create=mock_provider(monkeypatch,[good,{}, {'text':', forecast'}])
    from openai import OpenAI
    translator=i18n.EnglishTranslator('fake','fake','https://invalid.example')
    result=translator._translate_fixed_fact_prose(source,masked,tokens,OpenAI(api_key='fake'))
    assert not i18n.translation_validation_errors(source,result)
    assert translator.fixed_fact_fallback_calls==3 and create.call_count==3
    instruction=create.call_args.kwargs['messages'][0]['content']
    assert 'mapping every exact prose span ID' not in instruction
    assert 'JSON object {"text"' in instruction


def test_expanded_exact_nonfactual_qualification_is_restored_without_api():
    masked,tokens=i18n.protected_translation_input('结论证据不足，收入100',financial_context=True)
    part=next(part for part,needed in i18n.fixed_fact_prose_plan(masked,tokens) if needed and '结论' in part)
    result,errors=i18n.validate_fixed_prose_span(part,'Conclusion: insufficient evidence, revenue',tokens)
    assert not errors and '__KEEP_A__' in result
    assert i18n.validate_fixed_prose_span(part,'Conclusion: verified, revenue',tokens)[1]==['qualification marker mismatch']
    assert i18n.validate_fixed_prose_span(part,'Conclusion __KEEP_A__ __KEEP_B__',tokens)[1]==['unexpected protected marker']


def test_fixed_prose_diagnostics_report_categories_without_source_text(caplog,monkeypatch):
    source='收入100，上期80'
    masked,tokens=i18n.protected_translation_input(source,financial_context=True)
    create=mock_provider(monkeypatch,[{'0':'Revenue 999','2':', prior'},{'0':'Revenue 999'}])
    from openai import OpenAI
    translator=i18n.EnglishTranslator('fake','fake','https://invalid.example')
    with pytest.raises(i18n.TranslationError):
        translator._translate_fixed_fact_prose(source,masked,tokens,OpenAI(api_key='fake'))
    assert create.call_count==2
    assert 'numeric literals changed' in caplog.text and 'source_chars=' in caplog.text
    assert 'Revenue 999' not in caplog.text and source not in caplog.text


def test_mixed_order_and_numeric_failures_keep_their_separate_validation_paths(monkeypatch):
    sources=[f'收入{100+i}，上期{80+i}' for i in range(16)]
    reversed_text='Revenue __KEEP_B__, prior __KEEP_A__'
    invalid_numeric='Revenue __KEEP_A__, prior __KEEP_B__, invented 999'
    first={str(i):(reversed_text if i<3 else invalid_numeric if i==3 else 'Revenue __KEEP_A__, prior __KEEP_B__') for i in range(16)}
    second={str(i):(reversed_text if i<3 else invalid_numeric) for i in range(4)}
    create=mock_provider(monkeypatch,[first,second]+[{'0':'Revenue','2':', prior'}]*3+[{'text':'Revenue __KEEP_A__, prior __KEEP_B__'}])
    translator=i18n.EnglishTranslator('fake','fake','https://invalid.example')
    translator.prefetch(sources)
    assert all(not i18n.translation_validation_errors(source,translator.cache[source]) for source in sources)
    assert translator.fixed_fact_fallback_calls==3 and create.call_count==6
    assert all('999' not in translator.cache[source] for source in sources)


@pytest.mark.parametrize('source,expected',[
    ('扣非归母453.90亿元','net profit attributable to owners of the parent excluding non-recurring items'),
    ('扣非净利润453.90亿元','net profit excluding non-recurring items'),
    ('A股法定扣非','statutory profit excluding non-recurring items'),
    ('扣除非经常性损益后的归母净利润','net profit attributable to owners of the parent excluding non-recurring items'),
    ('扣非歸母淨利潤','net profit attributable to owners of the parent excluding non-recurring items'),
    ('扣除非經常性損益後的淨利潤','net profit excluding non-recurring items'),
    ('非经常性损益剔除后的净利润','net profit excluding non-recurring items'),
])
def test_non_recurring_exclusion_terms_are_immutable(source,expected):
    masked,tokens=i18n.protected_translation_input(source,financial_context=True)
    assert expected in tokens.values()
    result=i18n.restore_translation_tokens(masked.replace('A股','A-share '),tokens)
    assert not i18n.translation_validation_errors(source,result)
    bad=result.replace('excluding non-recurring items','').replace('net profit','non-recurring net profit')
    assert 'non-recurring exclusion meaning changed' in i18n.translation_validation_errors(source,bad)


def test_terminology_revision_invalidates_old_canonical_export_cache(tmp_path,monkeypatch):
    i18n.configure('en')
    source=tmp_path/'扣非报告.md';source.write_text('扣非净利润',encoding='utf-8')
    old_digest=hashlib.sha256(b'v81-financial-context-5\0'+source.read_bytes()).hexdigest()[:20]
    directory=tmp_path/'.english';directory.mkdir()
    old=directory/f'report_{old_digest}_en.md';old.write_text('Non-recurring net profit',encoding='utf-8')
    output=i18n.localize_artifact(source)
    assert output!=old and 'excluding non-recurring items' in output.read_text(encoding='utf-8')
    assert old.read_text(encoding='utf-8')=='Non-recurring net profit'


@pytest.mark.parametrize('unit',['个基点','個基點','个百分点','個百分點'])
def test_counted_financial_unit_is_one_immutable_fact_without_classifier_span(unit):
    source=f'净息差下降20{unit}，上期下降10{unit}'
    masked,tokens=i18n.protected_translation_input(source,financial_context=True)
    assert sum(literal in ('basis points','percentage points') for literal in tokens.values())==2
    assert all(part not in ('个','個') for part,needed in i18n.fixed_fact_prose_plan(masked,tokens) if needed)
    rendered=i18n.restore_translation_tokens(masked.replace('净息差下降','Net interest margin fell ').replace('，上期下降',', prior decline '),tokens)
    assert not i18n.translation_validation_errors(source,rendered)


@pytest.mark.parametrize('term',['普通股一级资本比率','普通股一級資本比率','核心一级资本比率'])
def test_common_equity_tier_one_term_never_adds_a_numeric_fact(term):
    source=term+'从14.6%升至14.8%'
    masked,tokens=i18n.protected_translation_input(source,financial_context=True)
    assert 'common equity tier one capital ratio' in tokens.values()
    assert all('一级' not in part and '一級' not in part for part,needed in i18n.fixed_fact_prose_plan(masked,tokens) if needed)
    rendered=i18n.restore_translation_tokens(masked.replace('从',' moved from ').replace('升至',' to '),tokens)
    assert not i18n.translation_validation_errors(source,rendered)


def test_fixed_prose_payload_hides_factual_literals_without_changing_restoration(monkeypatch):
    source='收入123.45亿港元，上期98.76亿港元，见[本期PDF第5页]'
    masked,tokens=i18n.protected_translation_input(source,financial_context=True)
    requested={str(i):part for i,(part,needed) in enumerate(i18n.fixed_fact_prose_plan(masked,tokens)) if needed}
    translated={key:part.replace('收入','Revenue').replace('，上期',', prior').replace('，见',', see') for key,part in requested.items()}
    create=mock_provider(monkeypatch,[translated])
    from openai import OpenAI
    translator=i18n.EnglishTranslator('fake','fake','https://invalid.example')
    result=translator._translate_fixed_fact_prose(source,masked,tokens,OpenAI(api_key='fake'))
    payload=json.loads(create.call_args.kwargs['messages'][1]['content'])
    serialized=json.dumps(payload)
    assert '123.45' not in serialized and '98.76' not in serialized and 'HKD' not in serialized
    assert '[Current filing PDF p. 5]' not in serialized
    assert 'HKD 123.45 hundred million' in result
    assert '[Current filing PDF p. 5]' in result
    assert not i18n.translation_validation_errors(source,result)


@pytest.mark.parametrize('original,expected',[
    ('Profit acceleration may also partly came from other income of 6,059,485 thousand.',
     'Profit acceleration may also partly come from other income of 6,059,485 thousand.'),
    ("but the current period's not obtained itemized details.",
     'but itemized details were not obtained for the current period.'),
    ('Cash received from the sale of goods and provision of services in the same period 950.87 hundred million, compared with the previous year 870.94 hundred million increased by 9.18%[Current filing PDF p. 31]',
     'Cash received from the sale of goods and provision of services in the same period was 950.87 hundred million, compared with 870.94 hundred million in the previous year, an increase of 9.18%[Current filing PDF p. 31]'),
    ('Capital expenditure (management reporting basis) in the first half 46,583 million，2024 same period of the year 23,088 million[Current filing PDF p. 15]',
     'Capital expenditure (management reporting basis) in the first half was 46,583 million, compared with the 2024 first-half figure of 23,088 million[Current filing PDF p. 15]'),
])
def test_observed_english_grammar_repairs_keep_every_fact_and_qualification(original,expected):
    assert i18n.normalize_english_grammar(original)==expected
    assert not i18n.translation_validation_errors(original,expected)
    assert re.findall(i18n._FIGURE_PATTERN,original)==re.findall(i18n._FIGURE_PATTERN,expected)
    assert i18n.offline_english_text(original)==expected


def test_hsbc_pillar_two_a_term_does_not_invent_an_arabic_two():
    source='第二A支柱要求為風險加權資產的2.6%，其中1.5%須由CET1滿足'
    masked,tokens=i18n.protected_translation_input(source,financial_context=True)
    assert 'Pillar Two A requirements' in tokens.values()
    result=i18n.restore_translation_tokens(masked.replace('為風險加權資產的',' equal ').replace('，其中',', of which ').replace('須由',' must be met by ').replace('滿足',''),tokens)
    assert not i18n.translation_validation_errors(source,result)
    assert i18n.numeric_tokens(source)==i18n.numeric_tokens(result)


def test_hsbc_standalone_genitive_particles_do_not_need_model_translation(monkeypatch):
    source='普通股一级资本比率从2024年12月31日的14.6%升至2025年6月30日的14.8%'
    masked,tokens=i18n.protected_translation_input(source,financial_context=True)
    plan=i18n.fixed_fact_prose_plan(masked,tokens)
    assert all(part.strip()!='的' for part,needed in plan if needed)
    requested={str(i):part for i,(part,needed) in enumerate(plan) if needed}
    output={key:part.replace('从',' from ').replace('升至',' to ') for key,part in requested.items()}
    create=mock_provider(monkeypatch,[output])
    from openai import OpenAI
    translator=i18n.EnglishTranslator('fake','fake','https://invalid.example')
    result=translator._translate_fixed_fact_prose(source,masked,tokens,OpenAI(api_key='fake'))
    assert not i18n.translation_validation_errors(source,result)
    assert create.call_count==1
    for prohibited in ('未','不','无','無'):
        assert i18n.validate_fixed_prose_span(prohibited,'',tokens)[1]==['missing or non-text span']


def test_existing_english_docx_author_is_repaired_without_model_or_body_changes(tmp_path,monkeypatch):
    from docx import Document
    directory=tmp_path/'.english';directory.mkdir()
    path=directory/'report_0123456789abcdef0123_en.docx'
    doc=Document();doc.add_paragraph('Verified source facts: revenue HKD 100 million [S1].')
    doc.core_properties.author='AI 排雷框架 V3.0';doc.core_properties.last_modified_by='排雷系统'
    doc.save(path)
    before=[p.text for p in Document(path).paragraphs]
    monkeypatch.setattr('openai.OpenAI',Mock(side_effect=AssertionError('No model calls')))
    i18n.configure('en');i18n.localize_artifact(path)
    repaired=Document(path)
    assert [p.text for p in repaired.paragraphs]==before
    assert repaired.core_properties.author==i18n.ENGLISH_SYSTEM_NAME
    assert repaired.core_properties.last_modified_by==i18n.ENGLISH_SYSTEM_NAME
    assert path.with_suffix('.docx.before_v81_english_repair').exists()


def test_new_english_docx_sets_english_creator_metadata_before_cache_validation(tmp_path,monkeypatch):
    from docx import Document
    path=tmp_path/'风险报告.docx'
    doc=Document();doc.add_paragraph('Evidence required.');doc.core_properties.author='AI 排雷框架 V3.0';doc.save(path)
    monkeypatch.setattr('openai.OpenAI',Mock(side_effect=AssertionError('No model calls')))
    i18n.configure('en');output=i18n.localize_artifact(path);converted=Document(output)
    assert converted.core_properties.author==i18n.ENGLISH_SYSTEM_NAME
    assert converted.core_properties.last_modified_by==i18n.ENGLISH_SYSTEM_NAME
    assert not i18n.HAN.search(' '.join((converted.core_properties.title,converted.core_properties.subject,converted.core_properties.author)))


@pytest.mark.parametrize('source',[
    '金額保留原幣單位，不能以此等式直接判斷可用現金；',
    '金额保留原币单位，不能以此等式直接判断可用现金；',
])
def test_generic_original_currency_unit_note_is_not_an_extra_financial_unit(source,monkeypatch):
    monkeypatch.setattr('openai.OpenAI',Mock(side_effect=AssertionError('No model calls')))
    translator=i18n.EnglishTranslator('fake','fake','https://invalid.example')
    result=translator.text(source)
    assert 'original reporting currency and unit of measure' in result
    assert 'cannot determine available cash' in result
    assert not i18n.translation_validation_errors(source,result)
    assert not i18n.financial_context_tokens(source) and not i18n.financial_context_tokens(result)
    translator.prefetch([source])
    assert translator.cache[source]==result


@pytest.mark.parametrize('source,term',[
    ('企業及機構理財業務費用及其他收益','Corporate and Institutional Banking'),
    ('合約服務收益撥回增加','release of contractual service margin'),
    ('新承保業務合約服務收益','new-business contractual service margin'),
    ('1年以內到期','maturing within 1 year'),
    ('全年ECL準備佔貸款總額平均值','full-year ECL charges as a proportion of average gross loans'),
    ('未填报','Not provided'),
    ('本期不再重现','did not recur in the current period'),
])
def test_confirmed_hsbc_financial_terms_are_locked(source,term):
    masked,tokens=i18n.protected_translation_input(source,financial_context=True)
    assert term in tokens.values()
    assert i18n.numeric_tokens(source)==i18n.numeric_tokens(i18n.restore_translation_tokens(masked,tokens))


def test_hsbc_semantic_migration_preserves_fact_order_period_and_allowance_balances():
    original=('HSBC. corporate and institutional wealth business fees increased 18%. '
              'an increase in the reversal of contract service revenue. '
              'Contract service revenue from new insurance business written for insurance products increased by USD 7 hundred million to USD 21 hundred million. '
              'of which 1 maturing within years USD166.59hundred million [Current filing PDF p. 66]. '
              'for full-year 2025, the ECL allowance will be approximately 40 basis points of average total loans. '
              'the full-year 2025 ECL provisions averaging approximately 40 basis points of total loans may be achieved. '
              'ECL allowance was USD 105.54 hundred million; ECL allowances were 1,725 million. '
              '2024 generated USD 36 hundred million net gains that will not recur. File number: No filing required.')
    migrated=i18n.offline_english_text(original)
    assert 'Corporate and Institutional Banking' in migrated
    assert 'release of contractual service margin' in migrated
    assert 'new-business contractual service margin' in migrated
    assert 'loans maturing within 1 year total USD 166.59 hundred million' in migrated
    assert 'ECL charges will be approximately 40 basis points of average gross loans' in migrated
    assert 'ECL charges equivalent to approximately 40 basis points of average gross loans may be achieved' in migrated
    assert 'ECL allowance was USD 105.54 hundred million' in migrated
    assert 'ECL allowances were 1,725 million' in migrated
    assert 'net gains that did not recur in the current period' in migrated
    assert 'File number: Not provided' in migrated
    assert re.findall(i18n._FIGURE_PATTERN,original)==re.findall(i18n._FIGURE_PATTERN,migrated)
    assert not i18n.translation_validation_errors(original,migrated)
    assert migrated.count('may')==original.count('may')


def test_hsbc_revenue_decrease_uses_correct_article_and_preserves_amounts():
    original=('HSBC. reported revenue USD 341.22 hundred million, compared with '
              'USD 372.92 hundred million decreased by USD 31.70 hundred million.')
    migrated=i18n.offline_english_text(original)
    assert 'a decrease of USD 31.70 hundred million' in migrated
    assert 'an decrease' not in migrated
    assert re.findall(i18n._FIGURE_PATTERN,original)==re.findall(i18n._FIGURE_PATTERN,migrated)
    assert not i18n.translation_validation_errors(original,migrated)
    legacy=('HSBC. Reported revenue was USD 341.22 hundred million, compared with '
            'USD 372.92 hundred million, an decrease of USD 31.70 hundred million.')
    assert 'a decrease of USD 31.70 hundred million' in i18n.offline_english_text(legacy)


def test_hsbc_business_semantic_repairs_do_not_change_other_issuers():
    original='Other Company: corporate and institutional wealth business and contract service revenue; ECL allowance was USD 105.54 hundred million.'
    assert i18n.offline_english_text(original)==original


def test_chinese_cib_and_csm_semantics_are_checked_beyond_numbers():
    assert 'banking segment name changed' in i18n.translation_validation_errors('企業及機構理財','corporate and institutional wealth business')
    assert 'insurance service margin meaning changed' in i18n.translation_validation_errors('合約服務收益','contract service revenue')
    assert not i18n.translation_validation_errors('企業及機構理財','Corporate and Institutional Banking')
    assert not i18n.translation_validation_errors('合約服務收益','contractual service margin')


@pytest.mark.parametrize('original,expected',[
    ('capital expenditures (management reporting basis) in the first half 46,583 million，2024 same period of the year 23,088 million[Current filing PDF p. 15]',
     'Capital expenditures (management reporting basis) in the first half was 46,583 million, compared with the 2024 first-half figure of 23,088 million[Current filing PDF p. 15]'),
    ('the complete notes to the three statements and segment notes for not obtained were not obtained',
     'the complete notes to the three statements and segment notes were not obtained'),
    ('the complete notes and segment details for not obtained',
     'the complete notes and segment details were not obtained'),
    ('The Hong Kong stock report not disclosed the metric; this metric not applicable. It may leading to losses.',
     'The Hong Kong stock report does not disclose the metric; this metric is not applicable. It may lead to losses.'),
])
def test_final_read_grammar_repairs_retain_negative_and_uncertain_meaning(original,expected):
    migrated=i18n.offline_english_text(original)
    assert migrated==expected
    assert not i18n.translation_validation_errors(original,migrated)
    if 'not' in original: assert 'not' in migrated
    assert original.count('may')==migrated.count('may')


@pytest.mark.parametrize('suffix',['md','docx','json'])
def test_existing_hsbc_downloads_repair_in_place_without_api(suffix,tmp_path,monkeypatch):
    from docx import Document
    directory=tmp_path/'.english';directory.mkdir()
    path=directory/f'report_0123456789abcdef0123_en.{suffix}'
    title='HSBC Holdings plc'
    body='corporate and institutional wealth business; reversal of contract service revenue; of which 1 maturing within years USD166.59hundred million [Current filing PDF p. 66].'
    if suffix=='md':path.write_text('# '+title+'\n\n'+body,encoding='utf-8')
    elif suffix=='json':path.write_text(json.dumps({'company':title,'analysis':body}),encoding='utf-8')
    else:
        doc=Document();doc.add_heading(title,0);doc.add_paragraph(body);doc.save(path)
    original=path.read_bytes()
    monkeypatch.setattr('openai.OpenAI',Mock(side_effect=AssertionError('No API calls')))
    assert i18n.repair_english_cache(path)==path
    if suffix=='docx':text='\n'.join(p.text for p in Document(path).paragraphs)
    else:text=path.read_text(encoding='utf-8')
    assert 'Corporate and Institutional Banking' in text and 'contractual service margin' in text
    assert 'within 1 year total USD 166.59 hundred million' in text
    assert path.with_suffix(path.suffix+'.before_v81_english_repair').read_bytes()==original
    repaired=path.read_bytes();i18n.repair_english_cache(path);assert path.read_bytes()==repaired
